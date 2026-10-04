import json
import subprocess
import sys

import pymupdf
from openpyxl import load_workbook

from l2c.compare.outputs import (
    assign_to_shop_files,
    by_plan_sheet,
    comparison_files,
    counts,
    coverage_lines,
    safe_name,
)
from l2c.compare.run import run_comparison
from l2c.contract.io import write_bundle
from l2c.mock.generate import FILE_A, FILE_B, mock_project


def test_safe_name_is_filesystem_safe_and_unique_per_path():
    a = safe_name("DA/Colonnes/Sample COLONNES Part 1.pdf")
    b = safe_name("DA/Dalles/Sample COLONNES Part 1.pdf")
    assert a != b and "/" not in a and " " not in a and a.startswith("Sample_COLONNES_Part_1-")


def test_every_finding_lands_in_exactly_one_shop_file_or_unassigned():
    m = mock_project()
    findings = run_comparison(m.bundle)
    by_file, unassigned = assign_to_shop_files(findings, m.bundle)
    assert set(by_file) == {FILE_A, FILE_B}
    placed = [f.id for fs in by_file.values() for f in fs] + [f.id for f in unassigned]
    assert len(placed) == len(set(placed))
    missing = next(f for f in findings if f.status == "missing")
    assert missing in by_file[FILE_A]  # N2 is covered by PART1 (sorted first); not unassigned


def test_plan_only_findings_with_no_covering_shop_file_are_unassigned():
    m = mock_project()
    findings = run_comparison(m.bundle)
    trimmed = [e for e in m.bundle.elements if not (e.source == "shop" and e.level == "N2")]
    m.bundle.elements = trimmed
    findings = run_comparison(m.bundle)
    _, unassigned = assign_to_shop_files(findings, m.bundle)
    assert any(f.level == "N2" for f in unassigned)


def test_comparison_files_counts_add_up():
    m = mock_project()
    findings = run_comparison(m.bundle)
    comps, unassigned = comparison_files(findings, m.bundle)
    total = sum(sum(c.counts.values()) for c in comps) + len(unassigned)
    plan_only_self = [
        f for f in findings if f.shop_ref is None and f.check_type != "cross.plan_vs_shop"
    ]
    assert total + len(plan_only_self) == len(findings)
    assert counts(findings)["non_compliant"] == 3


def test_by_plan_sheet_groups_by_title_block_sheet():
    m = mock_project()
    sections = by_plan_sheet(run_comparison(m.bundle), m.bundle)
    assert {"S-517", "S-518"} <= set(sections)


def test_coverage_lines_state_what_is_not_covered():
    lines = coverage_lines(mock_project().bundle)
    assert any(line.startswith("colonne: covered") for line in lines)
    assert any(line.startswith("dalle: NOT COVERED") for line in lines)


def test_cli_writes_all_outputs_and_counts_agree(tmp_path):
    meta, out = tmp_path / "meta", tmp_path / "out"
    write_bundle(meta, mock_project().bundle)
    r = subprocess.run(
        [sys.executable, "-m", "l2c.compare", str(meta), "--out", str(out)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert r.returncode == 0, r.stderr
    assert r.stdout.startswith("findings=") and "K-4" not in r.stdout
    names = {p.name for p in (out / "comparison").iterdir()}
    assert (
        sum(n.endswith(".json") for n in names) == 2 and sum(n.endswith(".pdf") for n in names) == 2
    )
    findings = json.loads((out / "findings.json").read_text(encoding="utf-8"))
    per_shop = [
        json.loads(p.read_text(encoding="utf-8")) for p in (out / "comparison").glob("*.json")
    ]
    assert all("counts" in c and c["contract_version"] for c in per_shop)
    # xlsx rows == findings
    ws = load_workbook(out / "findings.xlsx").active
    assert ws.max_row == len(findings) + 1
    assert [c.value for c in ws[1]][:2] == ["Feuillet", "Localisation"]
    # PDF text contains the status counts and the coverage statement
    text = "".join(p.get_text() for p in pymupdf.open(out / "report" / "by_plan_sheet.pdf"))
    assert "Sheet S-517" in text and "NOT COVERED" in text and "CONFIDENTIAL" in text
    shop_pdf = next((out / "comparison").glob("*.pdf"))
    assert "Shop drawing comparison" in pymupdf.open(shop_pdf)[0].get_text()


def test_comparison_of_an_empty_project_still_writes_every_output(tmp_path):
    from l2c.contract.io import MetaBundle

    meta, out = tmp_path / "meta", tmp_path / "out"
    write_bundle(meta, MetaBundle(project="EMPTY"))
    r = subprocess.run(
        [sys.executable, "-m", "l2c.compare", str(meta), "--out", str(out)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert r.returncode == 0, r.stderr
    assert (out / "findings.json").read_text(encoding="utf-8").strip() == "[]"
    assert (out / "report" / "summary.pdf").is_file() and (out / "findings.xlsx").is_file()
    text = "".join(p.get_text() for p in pymupdf.open(out / "report" / "summary.pdf"))
    assert "NOT COVERED" in text  # coverage is stated even when nothing was found
